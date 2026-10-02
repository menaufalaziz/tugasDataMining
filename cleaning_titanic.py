import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FILE_CSV = "data/train.csv"
FILE_SEBELUM = "output/dataset_sebelum_cleaning.xlsx"
FILE_SESUDAH = "output/datsert_setelah_cleaning.xlsx"

KOLOM_ID = "PassengerId"
KOLOM_NOMINAL = ["Survived", "Pclass", "Sex", "Embarked"]
KOLOM_NUMERIK = ["Age", "Fare"]
KOLOM_DIPAKAI = [KOLOM_ID] + KOLOM_NOMINAL + KOLOM_NUMERIK

def log(pesan):
    print(f"  - {pesan}")

# 1. Baca Data 
def baca_data_csv(path):
    df = pd.read_csv(path)
    print(f"Data diabaca : {df.shape[0]} baris dan {df.shape[1]} kolom")
    return df

def pilih_kolom(df):
    df = df.copy() 
    dibuang = [k for k in df.columns if k not in KOLOM_DIPAKAI]
    log(f"kolom tidak dipakai: {dibuang}")
    return df[KOLOM_DIPAKAI]

# 2. small function untuk cleaning
def hapus_duplicated(df):
    df = df.copy()
    sebelum = len(df)
    df = df.drop_duplicates(subset=[KOLOM_ID]).drop_duplicates()
    log(f"Baris duplikat dihapus : {sebelum - len(df)}")
    return df

def standarkan_teks(df):
    df = df.copy()
    df["Sex"] = df["Sex"].str.strip().str.lower()
    df["Embarked"] = df["Embarked"].str.strip().str.lower()
    log("Teks kolom Sec dan Embarked dirapikan")
    return df

def ubah_tipe_data(df):
    df = df.copy()
    for kolom in KOLOM_NUMERIK:
        df[kolom] = pd.to_numeric(df[kolom], errors="coerce")
    for kolom in KOLOM_NOMINAL:
        df[kolom] = df[kolom].astype("object")
    log("Tipe data diseuaiaka")
    return df


def isi_nilai_kosong_numerik(df):
    df = df.copy()
    # Mengisis nilai kosong dengan median
    kosong_age = df["Age"].isna().sum()
    median_kelompok = df.groupby(["Pclass", "Sex"])["Age"].transform("median")
    df["Age"] = df["Age"].fillna(median_kelompok).fillna(df["Age"].median())


    # fare dengen median keseluruhan
    kosong_fare = df["Fare"].isna().sum()
    df["Fare"] = df["Fare"].fillna(df["Fare"].median())

    return df

def isi_nilai_kosong_nominal(df):
    df = df.copy()
    for kolom in KOLOM_NOMINAL:
        kosong = df[kolom].isna().sum()
        if kosong:
            modus = df[kolom].mode()[0]
            df[kolom] = df[kolom].fillna(modus)
            log(f"{kolom}: {kosong} nilai kosong diisi modus ('{modus}')")
    return df
 
 
def tangani_outlier(df, kolom="Fare"):
    """Batasi nilai ekstrem dengan aturan IQR (winsorizing)."""
    df = df.copy()
    q1, q3 = df[kolom].quantile([0.25, 0.75])
    iqr = q3 - q1
    batas_bawah, batas_atas = q1 - 1.5 * iqr, q3 + 1.5 * iqr
    jumlah = ((df[kolom] < batas_bawah) | (df[kolom] > batas_atas)).sum()
    df[kolom] = df[kolom].clip(lower=batas_bawah, upper=batas_atas)
    log(
        f"{kolom}: {jumlah} outlier dibatasi ke "
        f"[{batas_bawah:.2f}, {batas_atas:.2f}] (aturan IQR)"
    )
    return df
 
 
def bulatkan_numerik(df):
    df = df.copy()
    df["Age"] = df["Age"].round(1)
    df["Fare"] = df["Fare"].round(2)
    return df

def bersihkan_data(df):
    return (
        df.pipe(pilih_kolom)
        .pipe(hapus_duplicated)
        .pipe(standarkan_teks)
        .pipe(ubah_tipe_data)
        .pipe(isi_nilai_kosong_nominal)
        .pipe(isi_nilai_kosong_numerik)
        .pipe(tangani_outlier, kolom="Fare")
        .pipe(bulatkan_numerik)
        .reset_index(drop=True)
    )
 
 
def ringkasan_kualitas(df, judul):
    print(f"\n[{judul}]")
    print(f"  Jumlah baris : {len(df)}")
    print(f"  Jumlah kolom : {df.shape[1]}")
    print(f"  Duplikat     : {df.duplicated().sum()}")
    kosong = df.isna().sum()
    print("  Nilai kosong per kolom:")
    for nama, jumlah in kosong.items():
        print(f"    {nama:<12} {jumlah}")

def validasi(df):
    assert len(df) >= 100, "Baris kurang dari 100"
    assert df.shape[1] >= 5, "Kolom kurang dari 5"
    assert df.isna().sum().sum() == 0, "Masih ada nilai kosong"
    assert df.duplicated().sum() == 0, "Masih ada duplikat"
    assert (df["Age"] > 0).all(), "Ada Age yang tidak valid"
    assert (df["Fare"] >= 0).all(), "Ada Fare yang tidak valid"
    print("\nValidasi lolos: tidak ada nilai kosong atau duplikat.")
    return df
 
 
# Tahap 4: menyimpan ke Excel dengan tampilan rapi
def simpan_excel(df, path, nama_sheet="Data"):
    df.to_excel(path, index=False, sheet_name=nama_sheet)
 
    wb = load_workbook(path)
    ws = wb[nama_sheet]
    font_biasa = Font(name="Arial", size=10)
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    isi_header = PatternFill("solid", start_color="305496")
 
    for baris in ws.iter_rows():
        for sel in baris:
            sel.font = font_biasa
    for sel in ws[1]:
        sel.font = font_header
        sel.fill = isi_header
        sel.alignment = Alignment(horizontal="center")
 
    for i, kolom in enumerate(df.columns, start=1):
        lebar = max(len(str(kolom)), df[kolom].astype(str).str.len().max()) + 3
        ws.column_dimensions[get_column_letter(i)].width = min(lebar, 40)
 
    ws.freeze_panes = "A2"
    wb.save(path)
    print(f"File disimpan: {path}")
    return df
 
 
# Program utama
def main():
    mentah = baca_data_csv(FILE_CSV)
    ringkasan_kualitas(mentah, "SEBELUM cleaning")
    simpan_excel(mentah, FILE_SEBELUM, nama_sheet="Data mentah")
 
    print("\nProses cleaning:")
    bersih = bersihkan_data(mentah)
 
    ringkasan_kualitas(bersih, "SESUDAH cleaning")
    validasi(bersih)
    simpan_excel(bersih, FILE_SESUDAH, nama_sheet="Data bersih")
 
 
if __name__ == "__main__":
    main()
